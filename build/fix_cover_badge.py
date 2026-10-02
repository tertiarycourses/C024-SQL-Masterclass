"""Redraw the cover badge caption (raster PNG): FUNDAMENTALS / WSQ · BEGINNERS -> MASTERCLASS / C024 · HANDS-ON."""
import io
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

F = "courseware/SQL Masterclass-v1.pptx"
p = Presentation(F)
sh = [s for s in p.slides[0].shapes if s.shape_type == 13 and s.image.size == (760, 760)][0]
im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
px = im.load()
# sample colours: blue caption text and grey sub-caption (darkest pixel in each band)
def darkest(y0, y1):
    return min((px[x, y] for x in range(150, 610) for y in range(y0, y1) if px[x, y][3] > 200), key=lambda c: sum(c[:3]))
blue, grey = darkest(595, 640), darkest(650, 690)
bg = px[380, 580]
d = ImageDraw.Draw(im)
d.rectangle([60, 585, 700, 700], fill=bg)
bold = "/System/Library/Fonts/Supplemental/Arial Bold.ttf"
reg = "/System/Library/Fonts/Supplemental/Arial.ttf"
def centre(text, font, y, fill):
    w = d.textlength(text, font=font)
    d.text(((760 - w) / 2, y), text, font=font, fill=fill)
centre("MASTERCLASS", ImageFont.truetype(bold, 46), 592, blue)
centre("C024  ·  HANDS-ON", ImageFont.truetype(reg, 31), 650, grey)
out = io.BytesIO(); im.save(out, "PNG")
sh.part.related_part(sh._element.blip_rId)._blob = out.getvalue()
p.save(F)
im.save("/tmp/c024_badge_new.png")
print("badge redrawn", blue, grey, bg)
