"""C024 Learner Guide: WSQ TGS-2020505790 LG -> non-WSQ SQL Masterclass LG."""
import re
from docx import Document

SRC = "source-wsq/courseware/LG-SQL Fundamental for Beginners.docx"
OUT = "courseware/LG-SQL Masterclass.docx"
REPO = "https://github.com/tertiarycourses/C024-SQL-Masterclass"
d = Document(SRC)


def nodes(par):
    return [t for t in par._p.iter() if t.tag.endswith("}t") and t.text]


def set_text(par, text):
    ts = nodes(par)
    ts[0].text = text
    for t in ts[1:]:
        if par.text.startswith(text) is False and t.text.strip().isdigit() and "\t" in text:
            continue
        t.text = ""


def sub(par, old, new):
    """Replace inside a single text node when possible (keeps run/hyperlink formatting)."""
    for t in nodes(par):
        if old in t.text:
            t.text = t.text.replace(old, new)
            return True
    return False


def drop(par):
    par._p.getparent().remove(par._p)


EXACT = {
    "SQL Fundamental for Beginners": "SQL Masterclass",
    "TGS Ref No: TGS-2020505790": "Course Code: C024",
    "Version 14": "Version 1",
    "Skills Framework: Data Engineering (ICT-DIT-3005-1.1), ICT Skills Framework.": None,
    "https://lms-tms.tertiaryinfotech.com": REPO,
    "Continue with the recommended WSQ courses at https://www.tertiarycourses.com.sg.":
        "Continue with the recommended courses at https://www.tertiarycourses.com.sg.",
    "Complete the TRAQOM feedback survey and download your e-certificate from https://lms-tms.tertiaryinfotech.com after being assessed Competent.":
        "Complete the course feedback form and collect your Certificate of Achievement from the trainer.",
}
PARTIAL = [
    ("the WSQ course SQL Fundamental for Beginners (TGS-2020505790)", "the course SQL Masterclass (C024)"),
    (" The guide is an approved open-book reference during the final assessment.", ""),
    ("download the datasets folder from the course materials on https://lms-tms.tertiaryinfotech.com.",
     f"download the labs/_all folder from the course repository at {REPO}."),
    ("LMS course materials.", "course repository."),
    ("Download the course mock data from the LMS.", "Download the course mock data from the course repository."),
    ("courseware/labs/", "labs/"),
    ("You want the datasets folder", "You want the labs/_all folder"),
]

in_assess = False
for p in list(d.paragraphs):
    t = p.text
    if p.style.name == "Heading 1":
        in_assess = t.strip() == "How You Will Be Assessed"
    if in_assess or t.startswith("How You Will Be Assessed\t"):
        drop(p)
        continue
    if t in EXACT:
        if EXACT[t] is None:
            drop(p)
        elif not sub(p, t, EXACT[t]):
            set_text(p, EXACT[t])
        continue
    for old, new in PARTIAL:
        if old in t and not sub(p, old, new):
            set_text(p, t.replace(old, new))
            t = p.text
    t = p.text
    new = re.sub(r"\s*\(TSC mapping: [^)]*\)", "", t)
    new = re.sub(r"\s*\((?:[AK]\d+(?:, )?)+\)\.$", ".", new)
    if new != t:
        for pat in (r"\s*\(TSC mapping: [^)]*\)", r"\s*\((?:[AK]\d+(?:, )?)+\)(?=\.$)"):
            for tn in nodes(p):
                tn.text = re.sub(pat, "", tn.text)
        if p.text != new:
            set_text(p, new)

vc = d.tables[0]
for row in list(vc.rows)[2:]:
    row._tr.getparent().remove(row._tr)
for c, v in zip(vc.rows[1].cells, ["1.0", "2 October 2026",
                                   "Initial release of the SQL Masterclass learner guide: step-by-step instructions for "
                                   "all 11 hands-on labs, SQL quick reference and glossary.", "Dr. Alfred Ang"]):
    set_text(c.paragraphs[0], v)

d.save(OUT)
print("wrote", OUT)
