"""C024 Lesson Plan: WSQ TGS-2020505790 LP -> non-WSQ SQL Masterclass LP."""
import copy, re, sys
sys.path.insert(0, "build")
from docx import Document
from slidemap import remap

SRC = "source-wsq/courseware/LP-SQL Fundamental for Beginners.docx"
OUT = "courseware/LP-SQL Masterclass.docx"
OLD_T, NEW_T = "SQL Fundamental for Beginners", "SQL Masterclass"
d = Document(SRC)


def set_text(par, text):
    runs = [r for r in par.runs if r.text] or par.runs
    if not runs:  # TOC entries: text sits in hyperlink runs
        ts = [t for t in par._p.iter() if t.tag.endswith("}t") and t.text]
        ts[0].text = text
        for t in ts[1:]:
            if not t.text.strip().isdigit():
                t.text = ""
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def set_cell(cell, text):
    set_text(cell.paragraphs[0], text)


def drop(el):
    el.getparent().remove(el)


# ---------------------------------------------------------------- paragraphs
in_assess = False
for p in list(d.paragraphs):
    t = p.text
    if p.style.name == "Heading 1":
        in_assess = t.strip() == "Assessment"
    if in_assess or t.strip().startswith("Assessment\t"):
        drop(p._p)
        continue
    if t.strip() == OLD_T:
        set_text(p, NEW_T)
    elif t.startswith("TGS Ref No:"):
        set_text(p, "Course Code: C024")
    elif t.strip() == "Version 14":
        set_text(p, "Version 1")
    elif t.startswith("Lab Reference (aligned to the TSC)"):
        set_text(p, t.replace(" (aligned to the TSC)", ""))
    elif t.startswith("Total training time:"):
        set_text(p, "Total instructional time: 450 minutes (7.5 hours).")

# ------------------------------------------------------------ version record
vc = d.tables[0]
for row in list(vc.rows)[2:]:
    drop(row._tr)
for c, v in zip(vc.rows[1].cells, ["1.0", "2 October 2026",
                                   "Initial release of the SQL Masterclass lesson plan: 1 day, 7.5 instructional "
                                   "hours (9:30 am – 5:30 pm), 11 hands-on labs across 4 topics.",
                                   "Dr. Alfred Ang"]):
    set_cell(c, v)

# --------------------------------------------------------- course information
for row in list(d.tables[1].rows):
    k = row.cells[0].text.strip()
    if k == "Course Title":
        set_cell(row.cells[1], NEW_T)
    elif k == "WSQ Course Reference":
        set_cell(row.cells[0], "Course Code")
        set_cell(row.cells[1], "C024")
    elif k == "Skills Framework":
        drop(row._tr)
    elif k == "Duration":
        set_cell(row.cells[1], "1 day · 7.5 instructional hours")
    elif k == "Daily Timing":
        set_cell(row.cells[1], "9:30 am – 5:30 pm (30-minute lunch break)")

# ------------------------------------------------------------ schedule table
sched = d.tables[2]
rows = list(sched.rows)[1:]
by = lambda pfx: next(r for r in rows if r.cells[2].text.startswith(pfx))
welcome, r_found, r_lab3, r_crud = by("Welcome"), by("SQL Foundations"), by("Hands-on: Lab 3"), by("INSERT, UPDATE")
lunch, r_t3, r_t4, recap = by("Lunch break"), by("Topic 3"), by("Topic 4"), by("Course recap")
for r in rows:
    if r not in (welcome, r_found, r_lab3, r_crud, lunch, r_t3, r_t4, recap):
        drop(r._tr)                                   # tea breaks, briefing, PP, OQ
drop(lunch._tr)
r_lab3._tr.addnext(lunch._tr)                         # lunch now 13:00-13:30, after Lab 4

LUNCH_NEW = 65                                        # 1-based position of the moved lunch slide
TEXT = {
    id(welcome): "Welcome, course introduction, trainer and learner introductions, ground rules, the SG Mart "
                 f"course dataset, lesson plan and course outline. Slides 1–{remap(25)}",
    id(r_found): re.sub(r"Slides \d+–\d+$", f"Slides {remap(26)}–{remap(54)}", r_found.cells[2].text),
    id(r_lab3): re.sub(r"Slides \d+–\d+$", f"Slides {remap(55)}–{LUNCH_NEW - 1}", r_lab3.cells[2].text),
    id(r_crud): re.sub(r"Slides \d+–\d+$", f"Slides {LUNCH_NEW + 1}–{remap(80)}", r_crud.cells[2].text),
    id(r_t3): re.sub(r"Slides \d+–\d+$", f"Slides {remap(82)}–{remap(103)}", r_t3.cells[2].text),
    id(r_t4): re.sub(r"Slides \d+–\d+$", f"Slides {remap(104)}–{remap(123)}", r_t4.cells[2].text),
    id(recap): f"Course recap and summary, Q&A, recommended courses and course feedback. Slides {remap(124)}–119",
}
# WSQ teaching minutes 30/60/75/60/60/45/15 (345) rescaled proportionally onto 450
order = [welcome, r_found, r_lab3, lunch, r_crud, r_t3, r_t4, recap]
src = [30, 60, 75, None, 60, 60, 45, 15]
teach = [m for m in src if m]
mins = [round(m / sum(teach) * 450 / 5) * 5 for m in teach]
mins[mins.index(max(mins))] += 450 - sum(mins)
it = iter(mins)
hm = lambda m: f"{9 + (30 + m) // 60}:{(30 + m) % 60:02d}"
clock = 0
for r, s in zip(order, src):
    m = 30 if s is None else next(it)
    set_cell(r.cells[0], f"{hm(clock)}–{hm(clock + m)}")
    set_cell(r.cells[1], f"{m} min")
    if id(r) in TEXT:
        set_cell(r.cells[2], TEXT[id(r)])
    clock += m
assert hm(clock) == "17:30", hm(clock)

# ------------------------------------------- lab reference: drop TSC column
lab = d.tables[3]
for row in lab.rows:
    drop(row.cells[1]._tc)
grid = lab._tbl.tblGrid
drop(grid.gridCol_lst[1])
grid.gridCol_lst[1].w = 1828800 * 2
for row in lab.rows:
    row.cells[1].width = 1828800 * 2

d.save(OUT)
print("wrote", OUT, "mins", mins)
